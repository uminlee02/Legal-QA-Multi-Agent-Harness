"""
docx_mcp_server.py — 공식 법률 답변서 워드(.docx) 생성 MCP 서버 (FastMCP, stdio).

  도구: create_legal_document(question, gist, conclusion, citations, summary, out_path)
        → 섹션 구조(질의/요지/관련법령·해설/결론) + 인용별 검증 배지 .docx 생성.

  핵심: '조문 원문'은 법제처 실제 텍스트(citations[].article_text, 코드가 주입)를 그대로 쓴다.
        요지/해설/결론만 생성 모델 담당. 검증 배지는 lawcheck.verify 결과(status/content).

  단독 실행: python docx_mcp_server.py   (stdio MCP 서버)
  agent_pipeline.py 가 MCP 클라이언트로 이 도구를 호출(생성 에이전트 → 문서작성 MCP).
"""
import os
import sys
import json
import shutil
import tempfile
import subprocess
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from mcp.server.fastmcp import FastMCP

# 한글 폰트: docx 기본이 영문 폰트(Cambria)면 한글이 □(두부)로 깨진다.
# run/스타일에 ascii·hAnsi·eastAsia 폰트를 모두 한글 폰트로 지정해야 한글에 적용됨.
KFONT = os.environ.get("DOC_FONT", "맑은 고딕")     # Windows 한글 기본. LibreOffice PDF 는 NanumGothic 로 폴백.
_FONT_ATTRS = ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs")


def _kfont_run(r):
    """run 에 한글 폰트(ascii/hAnsi/eastAsia/cs) 강제."""
    r.font.name = KFONT
    rf = r._element.get_or_add_rPr().get_or_add_rFonts()
    for a in _FONT_ATTRS:
        rf.set(qn(a), KFONT)


def _apply_doc_fonts(doc):
    """기본 스타일(Normal/Title/Heading)의 폰트를 한글 폰트로 지정."""
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            st = doc.styles[name]
            st.font.name = KFONT
            rf = st.element.get_or_add_rPr().get_or_add_rFonts()
            for a in _FONT_ATTRS:
                rf.set(qn(a), KFONT)
        except Exception:
            continue


def _finalize_fonts(doc):
    """모든 run(제목·본문·표·해설 포함)에 한글 폰트 강제 — add_paragraph/add_heading 로 만든 run 까지."""
    for p in doc.paragraphs:
        for r in p.runs:
            _kfont_run(r)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        _kfont_run(r)

mcp = FastMCP("legal-docx")


def _docx_to_pdf(docx_path, outdir):
    """LibreOffice headless로 docx → pdf. 실패/미설치 시 None 반환(에러 로그)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("[docx-MCP] soffice 미설치 → PDF 생략", file=sys.stderr)
        return None
    profile = tempfile.mkdtemp(prefix="lo_")           # 호출별 프로필(동시 실행 충돌 방지)
    try:
        subprocess.run(
            [soffice, "-env:UserInstallation=file://" + profile, "--headless",
             "--convert-to", "pdf", "--outdir", outdir, docx_path],
            timeout=90, capture_output=True, check=True)
        pdf = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
        return pdf if os.path.exists(pdf) else None
    except Exception as e:
        print(f"[docx-MCP] PDF 변환 실패: {e}", file=sys.stderr)
        return None
    finally:
        shutil.rmtree(profile, ignore_errors=True)

GREEN = RGBColor(0x1A, 0x7F, 0x37)
RED = RGBColor(0xC0, 0x2B, 0x2B)
GRAY = RGBColor(0x66, 0x66, 0x66)
NAVY = RGBColor(0x1F, 0x38, 0x64)


def _exist_badge(status):
    return {"real": ("✓ 실존", GREEN), "fake": ("✗ 미존재", RED)}.get(status, ("⚠ 확인필요", GRAY))


def _content_badge(status, content):
    if status != "real":
        return ("—", GRAY)
    return {"match": ("✓ 내용일치", GREEN), "mismatch": ("✗ 내용불일치", RED)}.get(
        content, ("· 내용미확인", GRAY))


def _run(p, text, *, bold=False, italic=False, color=None, size=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    _kfont_run(r)                       # 한글 폰트
    if color is not None:
        r.font.color.rgb = color
    if size is not None:
        r.font.size = Pt(size)
    return r


@mcp.tool()
def create_legal_document(question: str, gist: str, conclusion: str,
                          citations: list[dict], summary: dict,
                          out_path: str = "", out_format: str = "both") -> str:
    """공식 법률 질의 답변서를 .docx (및 선택적으로 .pdf) 로 생성한다.

    gist       : 답변 요지(생성, 2~3줄)
    conclusion : 결론(생성, 종합)
    citations  : [{label, status, content, article_text(법제처 원문), explanation(생성 해설)}]
    summary    : {n_total, n_real, n_match, n_content, n_fake, iterations}
    out_format : "docx" | "pdf" | "both"(기본). pdf 는 LibreOffice headless 변환.
    반환       : JSON 문자열 {"docx": 절대경로, "pdf": 절대경로 또는 null}
    """
    doc = Document()
    _apply_doc_fonts(doc)                # 한글 폰트 기본 지정(스타일)

    title = doc.add_heading("법률 질의 답변서", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("【질의 내용】", level=1)
    doc.add_paragraph(question or "")

    doc.add_heading("【답변 요지】", level=1)
    doc.add_paragraph(gist or "(요지 없음)")

    doc.add_heading("【관련 법령 및 해설】", level=1)
    if citations:
        for cit in citations:
            head = doc.add_paragraph()
            _run(head, f"▸ {cit.get('label','')}", bold=True, color=NAVY, size=11)
            head.add_run("   ")
            et, ec = _exist_badge(cit.get("status"))
            _run(head, f"[{et}]", bold=True, color=ec, size=9)
            head.add_run(" ")
            ct, cc = _content_badge(cit.get("status"), cit.get("content"))
            _run(head, f"[{ct}]", bold=True, color=cc, size=9)

            src = doc.add_paragraph()
            src.paragraph_format.left_indent = Pt(16)
            _run(src, "· 조문 원문: ", bold=True, size=9)
            src_text = (cit.get("article_text") or "").strip()
            if not src_text:
                src_text = ("(법제처에 실존하지 않는 조문)" if cit.get("status") == "fake"
                            else "(법제처 확인 불가 — 법령명·항 재확인 필요)")
            _run(src, src_text, size=9, color=RGBColor(0x33, 0x33, 0x33))

            expl = doc.add_paragraph()
            expl.paragraph_format.left_indent = Pt(16)
            _run(expl, "· 해설: ", bold=True, size=10)
            _run(expl, (cit.get("explanation") or "(해설 없음)").strip(), size=10)
    else:
        doc.add_paragraph("(추출된 법조문 인용 없음)")

    doc.add_heading("【결론】", level=1)
    doc.add_paragraph(conclusion or "(결론 없음)")

    sep = doc.add_paragraph()
    sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(sep, "─" * 30, color=GRAY)

    s = summary or {}
    v = doc.add_paragraph()
    _run(v, f"인용 검증: {s.get('n_real',0)}/{s.get('n_total',0)} 법제처 확인  |  "
            f"내용 일치: {s.get('n_match',0)}/{s.get('n_content',0)}  |  "
            f"환각: {s.get('n_fake',0)}건", bold=True, size=10)
    m = doc.add_paragraph()
    _gen_label = os.environ.get("DOC_GEN_LABEL", "EXAONE-3.5-7.8B")
    _cross_label = os.environ.get("DOC_CROSS_LABEL", "Qwen2.5-7B")
    _run(m, f"자율 수정: {s.get('iterations',0)}회  |  "
            f"생성: {_gen_label} · 교차검증: {_cross_label} (온프레미스)",
         size=9, color=GRAY)
    f = doc.add_paragraph()
    _run(f, f"문서 생성: {datetime.now():%Y-%m-%d}  |  법제처 국가법령정보 Open API 근거 기반 자동 검증",
         size=9, color=GRAY)

    if not out_path:
        out_path = f"법률답변서_{datetime.now():%Y%m%d_%H%M%S}.docx"
    out_path = os.path.abspath(out_path)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    _finalize_fonts(doc)                 # 모든 run 에 한글 폰트 강제(제목·본문·해설)
    doc.save(out_path)

    pdf_path = None
    if out_format in ("pdf", "both"):
        pdf_path = _docx_to_pdf(out_path, os.path.dirname(out_path))
    return json.dumps({"docx": out_path, "pdf": pdf_path}, ensure_ascii=False)


if __name__ == "__main__":
    mcp.run(transport="stdio")
